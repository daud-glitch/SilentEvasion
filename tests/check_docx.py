"""
Offline checks for the Word build - no LaTeX, no Word, no network.

    python tests/check_docx.py

The .docx is a second rendering of the same sources as the PDF, so the ways it
can be wrong are: stale (built before the last edit), short (a section, table or
figure silently dropped), or dirty (a LaTeX command leaked into the text). Each
of those is checked against the sources rather than against the converter's own
log line.
"""
import glob
import os
import re
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from docx import Document                                   # noqa: E402

PAPER = os.path.join(ROOT, "paper")
SECTIONS = os.path.join(PAPER, "sections")
TABLES = os.path.join(PAPER, "tables")
FIGURES = os.path.join(PAPER, "figures")
DOCX = os.path.join(PAPER, "main.docx")

FAILS = []


def check(label, ok, detail=""):
    print(("  OK   " if ok else "  FAIL ") + label
          + ((" - " + detail) if detail and not ok else ""))
    if not ok:
        FAILS.append(label)


def paragraphs(doc):
    return [p.text for p in doc.paragraphs]


def body_text(doc):
    return "\n".join(paragraphs(doc))


def test_fresh():
    print("\n[1] the .docx reflects the current sources")
    check("paper/main.docx exists", os.path.exists(DOCX))
    if not os.path.exists(DOCX):
        return
    stamp = os.path.getmtime(DOCX)
    stale = [p for p in
             glob.glob(os.path.join(SECTIONS, "*.tex"))
             + glob.glob(os.path.join(TABLES, "*.tex"))
             + glob.glob(os.path.join(FIGURES, "*.png"))
             if os.path.getmtime(p) > stamp]
    check("no source is newer than the .docx", not stale,
          "%d newer, e.g. %s -> run: python src/make_docx.py"
          % (len(stale), [os.path.basename(p) for p in stale[:3]]))


def test_structure(doc):
    print("\n[2] sections appear in the order main.tex gives them")
    h1 = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    check("first heading is the Abstract", h1 and h1[0] == "Abstract", str(h1[:2]))
    check("last heading is References", h1 and h1[-1] == "References", str(h1[-1:]))
    numbered = [h for h in h1 if re.match(r"^\d+\.\s", h)]
    check("eleven numbered sections", len(numbered) == 11, str(len(numbered)))
    check("section numbers are 1..11 in order",
          [int(h.split(".")[0]) for h in numbered] == list(range(1, 12)),
          str([h.split(".")[0] for h in numbered]))
    # The order is the argument: the three attack sections must be read before
    # the discussion, which is what the PDF now does too.
    titles = [h.lower() for h in numbered]
    def idx(fragment):
        for i, t in enumerate(titles):
            if fragment in t:
                return i
        return -1
    check("saliency before transfer before defence",
          -1 < idx("saliency") < idx("transfer") < idx("defence"), str(titles))
    check("all three results sections precede the discussion",
          max(idx("saliency"), idx("transfer"), idx("defence")) < idx("discussion"),
          str(titles))
    check("conclusion is last before References", idx("conclusion") == 10, str(titles))
    h2 = [p.text for p in doc.paragraphs if p.style.name == "Heading 2"]
    check("subsections are numbered x.y", all(re.match(r"^\d+\.\d+\s", h) for h in h2),
          str(h2[:3]))


def test_tables(doc):
    print("\n[3] every generated table is an editable Word table")
    files = sorted(glob.glob(os.path.join(TABLES, "*.tex")))
    check("one Word table per generated table", len(doc.tables) == len(files),
          "%d tables vs %d files" % (len(doc.tables), len(files)))
    text = " ".join(body_text(doc).split())
    # Tables appear in document order (t1..t4, then the pairs table, then
    # saliency, transfer, defence), not filename order, so each Word table is
    # matched to its file by the caption printed immediately above it.
    order = []
    for p in doc.paragraphs:
        m = re.match(r"^Table (\d+)\. ", p.text)
        if m:
            order.append(int(m.group(1)))
    check("table captions are numbered 1..N in document order",
          order == list(range(1, len(doc.tables) + 1)), str(order))
    captions = [p.text for p in doc.paragraphs if re.match(r"^Table \d+\. ", p.text)]
    for path in files:
        body = open(path, encoding="utf-8").read()
        cap = re.search(r"\\caption\{", body)
        if not cap:
            continue
        from src.make_docx import take_group
        caption, _ = take_group(body, cap.end() - 1)
        head = " ".join(caption.split())[:40].replace("\\", "")
        match = [c for c in captions if head.split(" ")[0] and head[:25] in c]
        check("caption of %s is present" % os.path.basename(path), bool(match),
              "looked for: " + head[:40])
        if not match:
            continue
        idx = captions.index(match[0])
        rows = len(re.findall(r"\\\\\s*$", body, re.M))
        check("%s has a Word row for every LaTeX row"
              % os.path.basename(path), len(doc.tables[idx].rows) == rows,
              "%d vs %d" % (len(doc.tables[idx].rows), rows))


def figure_sources():
    """Every .tex that can carry a \includegraphics, figure bodies included."""
    paths = sorted(glob.glob(os.path.join(SECTIONS, "*.tex")))
    paths += sorted(glob.glob(os.path.join(FIGURES, "*.tex")))
    out = []
    for path in paths:
        with open(path, encoding="utf-8") as handle:
            out.append(handle.read())
    return out


def test_figures(doc):
    # The count is read from the sources, not typed in here: a figure added to a
    # section must not need an edit to this test.
    wanted = sum(src.count("\\includegraphics") for src in figure_sources())
    print("\n[4] all %d figures are embedded images" % wanted)
    pngs = sorted(glob.glob(os.path.join(FIGURES, "*.png")))
    check("a PNG exists for every figure", len(pngs) == wanted,
          "%d PNG(s) for %d figure(s)" % (len(pngs), wanted))
    blips = [r for r in doc.part.rels.values()
             if "image" in r.reltype]
    check("one image is embedded per figure", len(blips) == wanted,
          "%d image(s) for %d figure(s)" % (len(blips), wanted))
    text = body_text(doc)
    missing = [i for i in range(1, wanted + 1) if ("Figure %d." % i) not in text]
    check("every figure caption is present", not missing, "missing %s" % missing)


def test_clean(doc):
    print("\n[5] no LaTeX survived the conversion")
    text = body_text(doc)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text += "\n" + cell.text
    leaks = sorted(set(re.findall(r"\\[a-zA-Z]+", text)))
    check("no backslash command in the text", not leaks, str(leaks[:6]))
    check("no unresolved reference", "\\ref" not in text and "??" not in text)
    check("no unresolved citation", "[None]" not in text and "\\cite" not in text)
    check("no stray braces from a failed group",
          "{" not in text and "}" not in text,
          str([s for s in re.findall(r".{20}[{}].{20}", text)][:2]))
    # Checked per paragraph: paragraphs are joined with a newline, and an empty
    # paragraph (a spacer after a table) would otherwise look like a double
    # space at a paragraph boundary.
    doubled = []
    for p in doc.paragraphs:
        if "  " in p.text:
            doubled += re.findall(r"\S{0,20}  \S{0,20}", p.text)
    check("no doubled spaces from a swallowed macro", not doubled,
          str(doubled[:3]))
    check("no percent sign was swallowed",
          text.count("%") >= 40, "%d percent signs" % text.count("%"))


def test_content(doc):
    print("\n[6] nothing was silently dropped")
    from src.make_docx import (body_order, collect_blocks, inline_runs,
                               meta_from_main, number_blocks)
    meta = meta_from_main()
    blocks = collect_blocks(body_order(meta["body"]))
    labels, cites = number_blocks(blocks)
    text = " ".join(body_text(doc).split())
    missing = []
    for b in blocks:
        if b["kind"] not in ("section", "subsection", "paragraph", "p"):
            continue
        runs = inline_runs(b["text"], labels, cites, {}, [])
        plain = " ".join("".join(t for t, _ in runs).split())
        if not plain:
            continue
        probe = plain[:45]
        if probe not in text:
            missing.append(probe)
    check("every heading and paragraph reached the .docx", not missing,
          "%d missing, e.g. %s" % (len(missing), missing[:3]))
    check("every citation key became a number",
          all(n > 0 for n in cites.values()), str(list(cites.items())[:3]))
    refs = [p for p in doc.paragraphs if re.match(r"^\[\d+\]", p.text)]
    check("references list has one entry per cited key", len(refs) == len(cites),
          "%d entries vs %d keys" % (len(refs), len(cites)))
    check("the self-replay result is in the Word file",
          "2,079" in text, "the number the transfer section quotes")
    # The availability section promises a repository; a link that survives the
    # PDF but is dropped by the Word build is not an available artefact.
    check("the artefact URL survives into the Word file",
          "github.com/daud-glitch/SilentEvasion" in text, "the \\url in section 10")


def main():
    print("DOCX CHECK")
    test_fresh()
    if not os.path.exists(DOCX):
        print("\nRESULT: FAIL: 1 problem(s)")
        return 1
    doc = Document(DOCX)
    test_structure(doc)
    test_tables(doc)
    test_figures(doc)
    test_clean(doc)
    test_content(doc)
    print("\nRESULT: %s" % ("FAIL: %d problem(s)" % len(FAILS) if FAILS else "OK"))
    return 1 if FAILS else 0


if __name__ == "__main__":
    sys.exit(main())
